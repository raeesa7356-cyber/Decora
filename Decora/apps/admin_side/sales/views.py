from django.shortcuts import render
from django.contrib.admin.views.decorators import staff_member_required
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.http import HttpResponse
from datetime import timedelta, datetime
from decimal import Decimal
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io
from apps.user_side.accounts.models import WalletTransaction
from apps.admin_side.orders.models import Order, OrderItem
from apps.admin_side.coupons.models import CouponUsage
from django.core.paginator import Paginator

def _get_date_range(request):
    period = request.GET.get('period', 'weekly')   
    custom_from = request.GET.get('date_from')
    custom_to = request.GET.get('date_to')
    today = timezone.now().date()

    if custom_from and custom_to:
        try:
            start = datetime.strptime(custom_from, '%Y-%m-%d').date()
            end = datetime.strptime(custom_to, '%Y-%m-%d').date()
            return start, end, 'custom'
        except ValueError:
            pass

    if period == 'weekly':
        start = today - timedelta(days=6)   
    elif period == 'monthly':
        start = today.replace(day=1)
    elif period == 'yearly':
        start = today.replace(month=1, day=1)
    else:  
        start = today

    return start, today, period


def _build_sales_report(start, end):
    orders = Order.objects.filter(
        created_at__date__gte=start,
        created_at__date__lte=end,
        status__in=['ACTIVE', 'SHIPPED', 'OUT_FOR_DELIVERY', 'DELIVERED', 'PENDING']
    ).select_related('user')

    total_orders = orders.count()
    gross_revenue = orders.aggregate(t=Sum('total_amount'))['t'] or Decimal('0')
    total_discount = orders.aggregate(t=Sum('discount_amount'))['t'] or Decimal('0')

    coupon_deductions = (
        CouponUsage.objects.filter(
            order__created_at__date__gte=start,
            order__created_at__date__lte=end
        ).select_related('coupon', 'order', 'user')
        .values(
            'coupon__code',
            'coupon__discount_type',
            'coupon__discount_amount',
        ).annotate(usage_count=Count('id')))

    net_revenue = gross_revenue

    return {
        'orders': orders,
        'total_orders': total_orders,
        'gross_revenue': gross_revenue,
        'total_discount': total_discount,
        'net_revenue': net_revenue,
        'coupon_deductions': coupon_deductions,
    }

@staff_member_required
def wallet_transaction_list(request):
    transactions = (
        WalletTransaction.objects
        .select_related('wallet', 'wallet__user')
        .order_by('-created_at')
    )

    purpose = request.GET.get("purpose")
    transaction_type = request.GET.get("type")
    search = request.GET.get("search")

    if purpose:
        transactions = transactions.filter(purpose=purpose)
    if transaction_type:
        transactions = transactions.filter(transaction_type=transaction_type)
    if search:
        transactions = transactions.filter(
                Q(wallet__user__first_name__icontains=search) |
                Q(wallet__user__last_name__icontains=search) |
                Q(wallet__user__email__icontains=search)
            )

    paginator = Paginator(transactions, 5)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'admin_side/sales/transactions.html', {
            'transactions': page_obj,'purpose': purpose or '',
            'transaction_type': transaction_type or '','search': search or '',
        })

@staff_member_required
def sales_report(request):
    start, end, period = _get_date_range(request)
    report = _build_sales_report(start, end)
    context = {
        **report,
        'start': start,'end': end,
        'period': period,'date_from': request.GET.get('date_from') or str(start),
        'date_to': request.GET.get('date_to') or str(end),
    }
    return render(request, 'admin_side/sales/sales_report.html', context)

@staff_member_required
def download_sales_pdf(request):
    start, end, period = _get_date_range(request)
    report = _build_sales_report(start, end)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []
    elements.append(Paragraph(f"Decora Sales Report", styles['Title']))
    elements.append(Paragraph(f"Period: {start} to {end}", styles['Normal']))
    elements.append(Spacer(1, 20))
    summary_data = [
    ['Metric', 'Value'],
    ['Total Orders', str(report['total_orders'])],
    ['Gross Revenue', f"₹{report['gross_revenue']:,.2f}"],
    ['Total Discounts', f"₹{report['total_discount']:,.2f}"],
    ['Net Revenue', f"₹{report['net_revenue']:,.2f}"],
    ]
    summary_table = Table(summary_data, colWidths=[250, 200])
    summary_table.setStyle(TableStyle([
    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D4AF37')),
    ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#1a1a1a'), colors.HexColor('#111111')]),
    ('TEXTCOLOR', (0, 1), (-1, -1), colors.white),
    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#333333')),
    ('FONTSIZE', (0, 0), (-1, -1), 10),
    ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph("Order Details", styles['Heading2']))
    elements.append(Spacer(1, 10))
    order_data = [['Order ID', 'Date', 'Customer', 'Amount', 'Discount', 'Status']]
    
    for order in report['orders'][:50]:  
        order_data.append([
            order.order_id,
            order.created_at.strftime('%d %b %Y'),
            order.user.get_full_name() or order.user.email,
            f"₹{order.total_amount:,.2f}",
            f"₹{order.discount_amount:,.2f}",
            order.get_status_display(),
        ])

    order_table = Table(order_data, colWidths=[90, 70, 110, 70, 70, 60])
    order_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D4AF37')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(order_table)
    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="sales_report_{start}_{end}.pdf"'
    return response

@staff_member_required
def download_sales_excel(request):
    start, end, period = _get_date_range(request)
    report = _build_sales_report(start, end)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    gold_fill = PatternFill("solid", fgColor="D4AF37")
    bold_font = Font(bold=True)
    header_font = Font(bold=True, color="000000")
    ws1.append(["Decora Sales Report"])
    ws1['A1'].font = Font(bold=True, size=14)
    ws1.append([f"Period: {start} to {end}"])
    ws1.append([])
    ws1.append(["Metric", "Value"])
    
    for cell in ws1[4]:
        cell.fill = gold_fill
        cell.font = header_font

    ws1.append(["Total Orders", report['total_orders']])
    ws1.append(["Gross Revenue", float(report['gross_revenue'])])
    ws1.append(["Total Discounts", float(report['total_discount'])])
    ws1.append(["Net Revenue", float(report['net_revenue'])])
    ws1.column_dimensions['A'].width = 25
    ws1.column_dimensions['B'].width = 20
    ws2 = wb.create_sheet("Orders")
    headers = ["Order ID", "Date", "Customer", "Email", "Amount", "Discount", "Status"]
    ws2.append(headers)
    
    for cell in ws2[1]:
        cell.fill = gold_fill
        cell.font = header_font

    for order in report['orders']:
        ws2.append([
            order.order_id,
            order.created_at.strftime('%d %b %Y %H:%M'),
            order.user.get_full_name() or order.user.username,
            order.user.email,
            float(order.total_amount),
            float(order.discount_amount),
            order.get_status_display(),
        ])

    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws2.column_dimensions[col].width = 20

    ws3 = wb.create_sheet("Coupon Usage")
    ws3.append(["Coupon Code", "Discount Type", "Discount Amount", "Times Used"])
    
    for cell in ws3[1]:
        cell.fill = gold_fill
        cell.font = header_font

    for c in report['coupon_deductions']:
        ws3.append([
            c['coupon__code'],
            c['coupon__discount_type'],
            float(c['coupon__discount_amount']),
            c['usage_count'],
        ])

    for col in ['A', 'B', 'C', 'D']:
        ws3.column_dimensions[col].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
    response['Content-Disposition'] = f'attachment; filename="sales_report_{start}_{end}.xlsx"'
    return response